import io
from typing import Tuple
import pdfplumber
import docx


def extract_text_from_file(file_bytes: bytes, filename: str) -> Tuple[str, str]:
    """
    Extracts text from PDF, DOCX, or TXT files.
    Returns: (text, file_type)
    Raises ValueError if unsupported or if PDF has no extractable text layer.
    """
    lower_name = filename.lower()
    
    if lower_name.endswith(".txt"):
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="replace")
        if not text.strip():
            raise ValueError("The uploaded TXT file is empty.")
        return text, "txt"

    elif lower_name.endswith(".docx"):
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text from tables if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)
            if not text.strip():
                raise ValueError("The uploaded DOCX file does not contain any readable text.")
            return text, "docx"
        except Exception as e:
            if "readable text" in str(e):
                raise
            raise ValueError(f"Failed to process DOCX file: {str(e)}")

    elif lower_name.endswith(".pdf"):
        try:
            extracted_pages = []
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page_idx, page in enumerate(pdf.pages):
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        extracted_pages.append(page_text.strip())

            if not extracted_pages:
                raise ValueError(
                    "The uploaded PDF has no extractable text layer. "
                    "Please upload a text-based PDF, DOCX, or TXT file."
                )

            full_text = "\n\n".join(extracted_pages)
            if not full_text.strip():
                raise ValueError(
                    "The uploaded PDF has no extractable text layer. "
                    "Please upload a text-based PDF, DOCX, or TXT file."
                )
            return full_text, "pdf"
        except ValueError:
            raise
        except Exception:
            raise ValueError(
                "The uploaded PDF has no extractable text layer. "
                "Please upload a text-based PDF, DOCX, or TXT file."
            )

    else:
        raise ValueError(
            f"Unsupported file format '{filename}'. Please upload a PDF, DOCX, or TXT contract."
        )
